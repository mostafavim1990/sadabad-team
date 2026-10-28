#!/usr/bin/env python3
"""
استخراج اطلاعات لازم از فایل اکسل پروژه و ساخت data.json برای داشبورد وب
فقط دو شیت خوانده می‌شود: «شیت اول» و «آخرین شیت» (شیت‌های میانی نادیده گرفته می‌شوند)

استفاده:
    python extract_data.py                 (اولین فایل xlsm/xlsx داخل همین پوشه)
    python extract_data.py Besharati.xlsm
"""
import sys, json
from pathlib import Path
from datetime import datetime, date, timezone

try:
    import openpyxl
except ImportError:
    print("نصب کنید:  pip install openpyxl"); sys.exit(1)

def norm(s):
    return str(s or "").replace("ي", "ی").replace("ك", "ک").replace("\u200c", "").replace(" ", "").strip()

def val(v):
    if v is None: return None
    if isinstance(v, (datetime, date)): return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer(): return int(v)
    if isinstance(v, str): return v.strip() or None
    return v

def num(v):
    return v if isinstance(v, (int, float)) else 0

# نام گزینه‌های منوی کشویی شیت اول که با نام آیتم‌های شیت آخر فرق دارند
ALIASES = {"آرماتور": "آرماتوربندی", "لوله کشی": "لوله کشی آب و فاضلاب", "پنجره": "پنجره UPVC",
           "در": "در داخلی", "کاشی سرامیک": "کاشی و سرامیک", "سنگکار": "سنگ کار",
           "گچکار": "گچ کار", "نزول": "وجه الالتزام"}

def find_labeled(ws, labels, label_col=10, value_col=11, max_row=60):
    """مقدار کنار اولین برچسبی که پیدا شود (برچسب‌ها به ترتیب اولویت)"""
    if isinstance(labels, str): labels = [labels]
    for label in labels:
        for r in range(1, max_row + 1):
            if norm(ws.cell(r, label_col).value) == norm(label):
                return ws.cell(r, value_col).value
    return None

def main():
    here = Path(__file__).resolve().parent
    if len(sys.argv) > 1:
        path = Path(sys.argv[1]);  path = path if path.is_absolute() else here / path
    else:
        found = sorted(list(here.glob("*.xlsm")) + list(here.glob("*.xlsx")))
        found = [p for p in found if not p.name.startswith("~$")]
        if not found: print("فایل اکسل پیدا نشد."); sys.exit(1)
        path = found[0]
    print("در حال خواندن:", path.name)
    wb = openpyxl.load_workbook(path, data_only=True)
    first, last = wb.worksheets[0], wb.worksheets[-1]

    # ---------- شیت اول: ردیف‌های هزینه (ستون B, C, F, H)
    hdr = next((r for r in range(1, 25) if "ردیف" in str(first.cell(r, 1).value or "")), 2)
    headers = [val(first.cell(hdr, c).value) for c in (2, 3, 6, 8)]
    rows, by_cat = [], {}
    for r in range(hdr + 1, (first.max_row or hdr) + 1):
        b, c, f, h = (val(first.cell(r, k).value) for k in (2, 3, 6, 8))
        if first.cell(r, 1).value not in (None, "") and (b or c):  # ردیف می‌تواند مثل 261/1 هم باشد
            rows.append([b, c, f if f != "-" else None, h])
            if b: by_cat[norm(b)] = by_cat.get(norm(b), 0) + num(h)

    spent = num(find_labeled(first, ["جمع کل", "جمع هزینه های انجام شده"])) or sum(num(x[3]) for x in rows)
    total = num(find_labeled(first, ["هزینه تخمین کل", "تخمین هزینه کل", "تخمین هزینه"]))
    remaining = find_labeled(first, ["مانده هزینه", "باقیمانده هزینه", "هزینه باقی مانده"])
    remaining = num(remaining) if remaining is not None else total - spent
    meters = num(find_labeled(first, "متراژ کل"))
    per_m = num(find_labeled(first, "هزینه ساخت به ازای هر متر مربع"))

    # ---------- شیت آخر: آیتم‌ها E3:E63
    title_date, title_text, title_proj = (val(last.cell(1, c).value) for c in (3, 4, 5))
    ph = [val(last.cell(2, c).value) for c in (5, 4, 3, 2, 1)]   # چینش معکوس: E D C B A
    items, summary, cats, used = [], [], [], set()
    for r in range(3, last.max_row + 1):
        e, d = val(last.cell(r, 5).value), val(last.cell(r, 4).value)
        a, b, c = (last.cell(r, k).value for k in (1, 2, 3))
        if r <= 63 and e:
            items.append([e, d, val(c), val(b), val(a)])
            key = norm(e); used.add(key)
            s = by_cat.get(key)
            if s is None:  # از نام مستعار منوی کشویی
                s = sum(v for k, v in by_cat.items() if norm(ALIASES.get(k, "")) == key or ALIASES.get(k) == e)
                for k in by_cat:
                    if norm(ALIASES.get(k, k)) == key: used.add(k)
            else: used.add(key)
            cats.append({"name": e, "sum": s})
        elif r > 63 and d:
            summary.append({"label": d, "values": [val(c), val(b), val(a)]})   # C, B, A
    # دسته‌هایی که در شیت اول هست ولی در لیست آیتم‌ها نیست (برای صحیح بودن جمع‌ها)
    for k, v in by_cat.items():
        if k not in used and not any(norm(ALIASES.get(k, "")) == norm(c["name"]) for c in cats):
            cats.append({"name": k, "sum": v})
    for c in cats: c["percent"] = (c["sum"] / spent * 100) if spent else 0

    data = {
        "project": val(first["D1"].value) or title_proj or path.stem,
        "sourceFile": path.name,
        "generatedAt": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S+00:00"),
        "summary": {"total": total, "spent": spent, "remaining": remaining,
                    "percent": (spent / total * 100) if total else 0,
                    "meters": meters, "costPerMeter": per_m},
        "categories": cats,
        "expenses": {"header": headers, "rows": rows},
        "pricing": {"date": title_date, "title": title_text, "project": title_proj,
                    "header": ph, "rows": items, "summary": summary},
    }
    with open(here / "data.json", "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False)
    print(f"data.json ساخته شد: {len(rows)} ردیف هزینه، {len(items)} آیتم تخمین، {len(cats)} دسته")
    print(f"  هزینه کل {total:,.0f} | انجام‌شده {spent:,.0f} | باقیمانده {remaining:,.0f}")

if __name__ == "__main__":
    main()
