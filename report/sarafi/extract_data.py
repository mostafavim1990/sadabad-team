#!/usr/bin/env python3
"""
استخراج اطلاعات چک‌ها از فایل اکسل صرافی و قرار دادن آن‌ها داخل خود صفحه index.html

استفاده (فایل اکسل، index.html و این اسکریپت در یک پوشه باشند):
    python extract_data.py
    python extract_data.py arz.xlsm

بعد از اجرا، index.html به‌روز می‌شود. همان را دوباره آپلود/باز کنید.
"""
import sys, re, json, os
from pathlib import Path
from datetime import datetime, date, timezone

try:
    import openpyxl
except ImportError:
    print("نصب کنید:  pip install openpyxl"); sys.exit(1)

START, END = "/*DATA_START*/", "/*DATA_END*/"

def text(v):
    if v is None: return ""
    if isinstance(v, (datetime, date)): return v.strftime("%Y/%m/%d")
    return str(v).strip()

def num(v):
    return v if isinstance(v, (int, float)) else 0

def norm(s):
    return text(s).replace("ي", "ی").replace("ك", "ک").replace("\u200c", "").replace(" ", "")

def main():
    here = Path(__file__).resolve().parent
    if len(sys.argv) > 1:
        path = Path(sys.argv[1]); path = path if path.is_absolute() else here / path
    else:
        found = [p for p in sorted(list(here.glob("*.xlsm")) + list(here.glob("*.xlsx"))) if not p.name.startswith("~$")]
        if not found: print("فایل اکسل پیدا نشد."); sys.exit(1)
        path = found[0]
    page = here / "index.html"
    if not page.exists(): print("index.html کنار این اسکریپت پیدا نشد."); sys.exit(1)

    print("در حال خواندن:", path.name)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]

    # سطر عنوان جدول (ستون A = «صاحب چک»)
    hdr = next((r for r in range(1, 30) if "صاحبچک" in norm(ws.cell(r, 1).value)), 3)
    checks = []
    for r in range(hdr + 1, (ws.max_row or hdr) + 1):
        a, b, c, d = (ws.cell(r, k).value for k in (1, 2, 3, 4))
        if a is None and b is None and c is None and d is None:
            continue
        checks.append({
            "owner": text(a), "amount": num(b), "from": text(c), "status": text(d),
            "state": text(ws.cell(r, 6).value),   # وضعیت چک (سررسید نشده / پاس شده / خرج شده)
            "due": text(ws.cell(r, 8).value),     # تاریخ سررسید (فقط برای رنگ‌بندی و فیلتر ماه)
        })

    # توان خرید: خانه‌ی کنار برچسب «مانده / توان خرید»
    power = 0
    for row in ws.iter_rows(min_row=1, max_row=60, min_col=1, max_col=30):
        for cell in row:
            if "توانخرید" in norm(cell.value):
                power = num(ws.cell(cell.row, cell.column + 1).value)
                break
        if power: break

    # خلاصه مالی: در فایل فعلی، عنوان‌ها در ستون M و مقدارها در ستون N هستند.
    summary = {"totalBuy": 0, "totalPaid": 0, "netProfit": 0, "marketCapital": 0}
    labels = {
        "جمعخرید": "totalBuy",
        "جمعپرداختی": "totalPaid",
        "سود(خالص)": "netProfit",
        "سودخالص": "netProfit",
        "سرمایهدربازار": "marketCapital",   # سرمایه در بازار
    }
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 60, 60), min_col=1, max_col=30):
        for cell in row:
            key = labels.get(norm(cell.value))
            if key:
                value = ws.cell(cell.row, cell.column + 1).value
                summary[key] = num(value)
    # اگر سلول‌های خلاصه در ستون‌های استاندارد باشند، مستقیماً M/N را هم بررسی کن.
    for r in range(1, 10):
        key = labels.get(norm(ws.cell(r, 13).value))
        if key:
            summary[key] = num(ws.cell(r, 14).value)

    # تاریخ به‌روزرسانی: زمان آخرین ذخیره‌ی فایل اکسل
    mod = wb.properties.modified
    if mod is None:
        mod = datetime.fromtimestamp(os.path.getmtime(path))
    updated = mod.replace(tzinfo=timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ") if mod.tzinfo is None else mod.isoformat()

    data = {"source": path.name, "updatedAt": updated, "buyingPower": power, **summary, "checks": checks}
    blob = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    html = page.read_text(encoding="utf-8")
    i, j = html.find(START), html.find(END)
    if i < 0 or j < 0: print("نشانه‌های داده داخل index.html پیدا نشد."); sys.exit(1)
    page.write_text(html[:i + len(START)] + "window.DATA=" + blob + ";" + html[j:], encoding="utf-8")

    print(f"index.html به‌روز شد: {len(checks)} چک | جمع مبلغ چک‌ها {sum(c['amount'] for c in checks):,.0f} | توان خرید {power:,.0f}")

if __name__ == "__main__":
    main()
