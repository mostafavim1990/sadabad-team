# -*- coding: utf-8 -*-
"""خواندن cheques.xlsm و ذخیره مستقیم اطلاعات داخل index.html؛ بدون GitHub و data.json."""
from pathlib import Path
from datetime import date
import json, re, traceback

def gregorian_to_jalali(gy, gm, gd):
    g_days=[31,28,31,30,31,30,31,31,30,31,30,31]
    j_days=[31,31,31,31,31,31,30,30,30,30,30,29]
    gy-=1600; gm-=1; gd-=1
    day_no=365*gy+(gy+3)//4-(gy+99)//100+(gy+399)//400
    for i in range(gm): day_no+=g_days[i]
    if gm>1 and ((gy%4==0 and gy%100!=0) or gy%400==0): day_no+=1
    day_no+=gd
    j_day_no=day_no-79
    j_np=j_day_no//12053; j_day_no%=12053
    jy=979+33*j_np+4*(j_day_no//1461); j_day_no%=1461
    if j_day_no>=366:
        jy+=(j_day_no-1)//365
        j_day_no=(j_day_no-1)%365
    i=0
    while i<11 and j_day_no>=j_days[i]:
        j_day_no-=j_days[i]; i+=1
    return jy, i+1, j_day_no+1

def normalize_digits(value):
    return str(value or "").strip().translate(str.maketrans(
        "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789"))

def parse_jalali(value):
    if value is None: return None
    match=re.search(r"(\d{4})\s*[/-]\s*(\d{1,2})\s*[/-]\s*(\d{1,2})",normalize_digits(value))
    if not match: return None
    try: return tuple(map(int,match.groups()))
    except ValueError: return None

def get_period(value, today):
    due=parse_jalali(value)
    if due is None: return "future"
    if due<today: return "past"
    if due[:2]==today[:2]: return "current"
    ny,nm=(today[0]+1,1) if today[1]==12 else (today[0],today[1]+1)
    if due[:2]==(ny,nm): return "next"
    return "future"

def clean(v):
    if v is None: return ""
    if isinstance(v,float) and v.is_integer(): return int(v)
    if isinstance(v,(int,float)): return v
    return str(v).strip()

def sort_date_key(value):
    parsed = parse_jalali(value)
    # تاریخ‌های نامشخص در انتهای فهرست قرار می‌گیرند.
    return parsed if parsed is not None else (9999, 99, 99)

def is_passed(status):
    text = normalize_digits(status).replace("‌", " ").replace("-", " ").strip()
    text = " ".join(text.split())
    return "پاس شده" in text or "وصول شده" in text or "نقد شده" in text

def main():
    try:
        import openpyxl
    except ImportError:
        print("کتابخانه openpyxl نصب نیست. در CMD اجرا کن: py -m pip install openpyxl")
        return False

    folder=Path(__file__).resolve().parent
    excel_path=folder/"cheques.xlsm"
    template_path=folder/"index_template.html"
    output_path=folder/"index.html"
    if not excel_path.exists():
        print(f"فایل اکسل پیدا نشد: {excel_path}")
        return False
    if not template_path.exists():
        print(f"قالب صفحه وب پیدا نشد: {template_path}")
        return False

    workbook=None
    try:
        print("در حال خواندن فایل اکسل...")
        workbook=openpyxl.load_workbook(excel_path,data_only=True,keep_vba=True,read_only=True)
        sheet=workbook[workbook.sheetnames[0]]
        now=date.today()
        today=gregorian_to_jalali(now.year,now.month,now.day)
        records=[]
        # ستون‌های B، C، F، H و I؛ داده‌ها از ردیف 5 شروع می‌شوند.
        for values in sheet.iter_rows(min_row=5,min_col=2,max_col=9,values_only=True):
            selected=[values[0],values[1],values[4],values[6],values[7]]
            if not any(v not in ("",None) for v in selected): continue
            project,recipient,status,due_date,amount=[clean(v) for v in selected]
            passed = is_passed(status)
            records.append({"project":project,"recipient":recipient,"status":status,
                            "dueDate":due_date,"amount":amount,"period":get_period(due_date,today),
                            "passed":passed,"_sortDate":sort_date_key(due_date)})
        # مرتب‌سازی بر اساس نزدیک‌ترین تاریخ سررسید؛ تاریخ‌های نامشخص در انتها.
        records.sort(key=lambda row: row["_sortDate"])
        for row in records:
            row.pop("_sortDate", None)
        workbook.close(); workbook=None
        payload={"updatedAt":f"{today[0]:04d}/{today[1]:02d}/{today[2]:02d}","rows":records}
        html=template_path.read_text(encoding="utf-8")
        if "__DATA_JSON__" not in html:
            raise ValueError("نشانگر __DATA_JSON__ در index_template.html پیدا نشد.")
        html=html.replace("__DATA_JSON__",json.dumps(payload,ensure_ascii=False,separators=(",",":")).replace("</","<\\/"))
        temp=output_path.with_suffix(".html.tmp")
        temp.write_text(html,encoding="utf-8")
        temp.replace(output_path)
        print(f"موفقیت‌آمیز بود: {len(records)} ردیف در index.html ذخیره شد.")
        print(f"اکنون این فایل را باز کن: {output_path}")
        print("هیچ اطلاعاتی به GitHub ارسال نشد و فایل اکسل تغییر نکرد.")
        return True
    except Exception:
        print("هنگام ساخت صفحه وب خطا رخ داد:")
        traceback.print_exc()
        return False
    finally:
        if workbook is not None:
            try: workbook.close()
            except Exception: pass

if __name__=="__main__":
    try:
        ok=main()
        if not ok: input("\nبرای بستن پنجره Enter را بزنید...")
    except Exception:
        traceback.print_exc()
        input("\nبرای بستن پنجره Enter را بزنید...")
