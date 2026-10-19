#!/usr/bin/env python3
"""
استخراج اطلاعات از فایل اکسل پروژه و ساخت data.json برای داشبورد وب

نحوه استفاده (در همان پوشه‌ای که اکسل و این اسکریپت هستند):
    python extract_data.py
    python extract_data.py Farhangian.xlsm
    python extract_data.py Chitsazi.xlsm

بعد از اجرا، data.json آپدیت می‌شود.
سپس در صفحه وب دکمه «بارگذاری مجدد اطلاعات» را بزنید.
"""

import sys
import json
from pathlib import Path
from datetime import datetime, date, timezone

try:
    import openpyxl
except ImportError:
    print("خطا: کتابخانه openpyxl نصب نیست.")
    print("نصب کنید با:  pip install openpyxl")
    sys.exit(1)


def cell_value(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def extract_sheet(ws):
    max_r = 0
    max_c = 0
    scan_rows = min(ws.max_row or 1, 800)
    scan_cols = min(ws.max_column or 1, 40)

    for row in ws.iter_rows(min_row=1, max_row=scan_rows, max_col=scan_cols):
        for cell in row:
            if cell.value is not None:
                max_r = max(max_r, cell.row)
                max_c = max(max_c, cell.column)

    if max_r == 0:
        return {
            "name": ws.title,
            "title": ws.title,
            "headerRow": None,
            "rows": [],
            "rowCount": 0,
        }

    header_row = None
    for r in range(1, min(25, max_r + 1)):
        for c in range(1, min(12, max_c + 1)):
            v = ws.cell(r, c).value
            if v and "ردیف" in str(v):
                header_row = r
                break
        if header_row:
            break

    title = ws.title
    for coord in ("D1", "A1", "B1", "C1", "A2"):
        v = ws[coord].value
        if v and str(v).strip():
            title = str(v).strip()
            break

    rows = []
    for r in range(1, max_r + 1):
        cells = []
        has_content = False
        for c in range(1, max_c + 1):
            val = cell_value(ws.cell(r, c).value)
            cells.append(val)
            if val is not None and val != "":
                has_content = True
        if has_content:
            while cells and cells[-1] is None:
                cells.pop()
            rows.append({"rowNumber": r, "cells": cells})

    return {
        "name": ws.title,
        "title": title,
        "headerRow": header_row,
        "rows": rows,
        "rowCount": len(rows),
    }


def extract_chart(wb):
    """پیدا کردن جمع کل، تخمین هزینه و باقیمانده از روی برچسب سلول‌ها."""
    spent = remaining = total = None

    def consider(label, value):
        nonlocal spent, remaining, total
        if not isinstance(label, str) or not isinstance(value, (int, float)):
            return
        text = label.replace("ي", "ی").replace("ك", "ک").strip()
        if text in ("جمع کل", "جمع هزینه های انجام شده") and spent is None:
            spent = value
        elif (
            text.startswith("تخمین هزینه") or text.startswith("هزینه تخمین")
        ) and "آپدیت" not in text and "به روز" not in text and total is None:
            total = value
        elif (
            text.startswith("هزینه باقی") or text.startswith("مانده هزینه")
        ) and remaining is None:
            remaining = value

    if "Sheet1" in wb.sheetnames:
        ws = wb["Sheet1"]
        for r in range(1, min(ws.max_row or 1, 40) + 1):
            for c in range(1, min(ws.max_column or 1, 20)):
                consider(ws.cell(r, c).value, ws.cell(r, c + 1).value)

    if (spent is None or total is None) and any(name.startswith("Sheet") for name in wb.sheetnames):
        for name in wb.sheetnames:
            ws = wb[name]
            for r in range(1, min(ws.max_row or 1, 120) + 1):
                label = ws.cell(r, 4).value
                if isinstance(label, str) and "جمع کل" in label:
                    if spent is None:
                        spent = ws.cell(r, 1).value
                    if total is None:
                        total = ws.cell(r, 2).value
                    break

    spent = float(spent or 0)
    total = float(total or 0)
    remaining = float(remaining if remaining is not None else (total - spent if total else 0))
    percent = (spent / total * 100) if total else 0
    return {
        "spent": spent,
        "remaining": remaining,
        "total": total,
        "percent": percent,
    }


def detect_project_name(wb, source_file: str) -> str:
    if "Sheet1" in wb.sheetnames:
        v = wb["Sheet1"]["D1"].value
        if v:
            name = str(v).strip()
            for prefix in ("پروژه خیابان رجایی - ", "پروژه "):
                if name.startswith(prefix):
                    name = name[len(prefix):].strip()
            if name:
                return name if name.startswith("پروژه") else f"پروژه {name}"

    stem = Path(source_file).stem
    mapping = {
        "Farhangian": "پروژه فرهنگیان",
        "Chitsazi": "پروژه چیت‌سازی",
        "chitsazi": "پروژه چیت‌سازی",
    }
    return mapping.get(stem, f"پروژه {stem}")


def main():
    script_dir = Path(__file__).resolve().parent

    if len(sys.argv) > 1:
        excel_path = Path(sys.argv[1])
        if not excel_path.is_absolute():
            excel_path = script_dir / excel_path
    else:
        candidates = [
            script_dir / "Farhangian.xlsm",
            script_dir / "Chitsazi.xlsm",
            script_dir / "farhangian.xlsm",
            script_dir / "chitsazi.xlsm",
        ]
        excel_path = next((p for p in candidates if p.exists()), None)
        if not excel_path:
            print("فایل اکسل پیدا نشد.")
            print("استفاده:  python extract_data.py نام‌فایل.xlsm")
            sys.exit(1)

    if not excel_path.exists():
        print(f"فایل پیدا نشد: {excel_path}")
        sys.exit(1)

    print(f"در حال خواندن: {excel_path.name} ...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    data = {
        "project": detect_project_name(wb, excel_path.name),
        "sourceFile": excel_path.name,
        "generatedAt": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S+00:00"),
        "chart": extract_chart(wb),
        "sheets": [extract_sheet(wb[name]) for name in wb.sheetnames],
    }

    out_path = script_dir / "data.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False)

    chart = data["chart"]
    print("data.json ساخته شد.")
    print(f"  پروژه     : {data['project']}")
    print(f"  هزینه کل  : {chart['total']:,.0f}")
    print(f"  انجام‌شده : {chart['spent']:,.0f}")
    print(f"  باقیمانده : {chart['remaining']:,.0f}")
    print(f"  درصد      : {chart['percent']:.1f}%")
    print()
    print("حالا صفحه وب را باز کنید و دکمه «بارگذاری مجدد اطلاعات» را بزنید.")


if __name__ == "__main__":
    main()
